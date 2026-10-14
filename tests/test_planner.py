import sys
from pathlib import Path

import numpy as np
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from planner import Config, fifo_lead_buckets, read_input, solve, write_report  # noqa: E402

NO_SAFETY = Config(safety_stock=0)


def balance_ok(res, demand, init=0.0):
    net = res.inventory - res.backlog
    prev = np.concatenate([np.full((demand.shape[0], 1), init), net[:, :-1]], axis=1)
    return np.allclose(net, prev + res.release - demand)


def test_produces_just_in_time_when_capacity_allows():
    demand = np.array([[0, 10, 0, 5]])
    res = solve(demand, np.array([100, 100, 100, 100]), NO_SAFETY)
    assert res.release.tolist() == [[0, 10, 0, 5]]
    assert res.backlog.sum() == 0 and res.inventory.sum() == 0


def test_builds_ahead_instead_of_backlog():
    # Day 3 needs 150 but capacity is 100/day: 50 must be made earlier.
    demand = np.array([[0, 0, 150]])
    res = solve(demand, np.array([100, 100, 100]), NO_SAFETY)
    assert res.backlog.sum() == 0
    assert res.release[0, 0] == 0  # not earlier than needed
    assert res.release[0].tolist() == [0, 50, 100]
    assert balance_ok(res, demand)


def test_backlog_when_capacity_is_short_and_capacity_respected():
    demand = np.array([[80, 0, 0], [80, 0, 0]])
    cap = np.array([100, 100, 100])
    res = solve(demand, cap, NO_SAFETY)
    assert (res.release.sum(axis=0) <= cap).all()
    assert res.backlog[:, 0].sum() == 60
    assert res.backlog[:, -1].sum() == 0
    assert (res.release >= 0).all()
    assert balance_ok(res, demand)


def test_safety_stock_is_held():
    demand = np.array([[10, 10, 10, 10]])
    res = solve(demand, np.array([100] * 4), Config(safety_stock=20))
    assert res.inventory[0].tolist() == [20, 20, 20, 20]


def test_fifo_buckets():
    release = np.array([[10, 0, 0, 5, 0]])
    demand = np.array([[0, 0, 10, 0, 8]])
    b = fifo_lead_buckets(release, demand)
    # 10 made day 0 for day 2 (2 days early); 5 made day 3 for day 4 (1 day early); 3 never made
    assert b.tolist() == [[0, 5, 10, 0, 3]]


def test_sample_input_end_to_end(tmp_path):
    inp = read_input(ROOT / "data" / "primer_ulaz.xlsx")
    assert inp.demand.shape == (15, 42)
    assert inp.week_of_day[-1] == 6
    res = solve(inp.demand, inp.capacity)
    assert (res.release.sum(axis=0) <= inp.capacity + 1e-6).all()
    assert balance_ok(res, inp.demand)
    buckets = fifo_lead_buckets(res.release, inp.demand)
    assert buckets.sum() == pytest.approx(inp.demand.sum())
    out = write_report(inp, res, buckets, tmp_path / "plan.xlsx")
    from openpyxl import load_workbook

    wb = load_workbook(out)
    assert wb.sheetnames == ["Plan", "Dnevni pregled", "Nedeljni pregled", "Backlog", "Proizvodnja unapred"]
