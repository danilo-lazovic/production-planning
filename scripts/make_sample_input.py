"""Generate data/primer_ulaz.xlsx: a synthetic input in the same layout as the
real planning sheet. All numbers are random; no real company data.

    python scripts/make_sample_input.py [--pns 15] [--weeks 6] [--seed 7]
"""

from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font

DAY_NAMES = ["Pon", "Uto", "Sre", "Cet", "Pet", "Sub", "Ned"]
OUT = Path(__file__).resolve().parents[1] / "data" / "primer_ulaz.xlsx"


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--pns", type=int, default=15)
    p.add_argument("--weeks", type=int, default=6)
    p.add_argument("--seed", type=int, default=7)
    p.add_argument("--out", type=Path, default=OUT)
    args = p.parse_args()

    rng = np.random.default_rng(args.seed)
    days = 7 * args.weeks
    weekday = np.arange(days) % 7
    capacity = np.select([weekday < 5, weekday == 5], [450, 300], 0)

    # Sparse orders per PN; the first week gets a demand peak above capacity, so the
    # plan starts with backlog that has to be worked off.
    demand = np.zeros((args.pns, days), dtype=int)
    for pn in range(args.pns):
        size = rng.integers(10, 60)
        has_order = (rng.random(days) < 0.45) & (weekday < 6)
        demand[pn] = np.where(has_order, rng.poisson(size, days), 0)
    peak = slice(1, 6)
    demand[:, peak] = (demand[:, peak] * 2.2).astype(int)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    bold = Font(bold=True)
    col0 = 3  # data starts in column C, as in the original sheet
    for d in range(days):
        c = col0 + d
        ws.cell(1, c, "Kapacitet").font = bold
        ws.cell(2, c, int(capacity[d]))
        if d % 7 == 0:
            ws.cell(8, c, f"CW{d // 7 + 1:02d}").font = bold
        ws.cell(9, c, DAY_NAMES[d % 7]).font = bold

    row = 11
    for pn in range(args.pns):
        ws.cell(row, 1, f"PN-{pn + 1:04d}").font = bold
        ws.cell(row + 1, 1, "Planned release")
        ws.cell(row + 2, 1, "Planed order")
        ws.cell(row + 3, 1, "Cummulative")
        for d in range(days):
            ws.cell(row + 2, col0 + d, int(demand[pn, d]))
        row += 4
    ws.column_dimensions["A"].width = 16

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"Saved {args.out}: {args.pns} PN, {days} days, "
          f"demand {demand.sum()} vs capacity {capacity.sum()}")


if __name__ == "__main__":
    main()
