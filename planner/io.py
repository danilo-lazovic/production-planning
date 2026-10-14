"""Reading the planning sheet and writing the report workbook.

Expected layout (first sheet), with data starting in column C:

    row with "Kapacitet" in column C   -> next row holds daily capacity
    row with "Pon", "Uto", ...         -> day names (one column per day)
    "PN-..." in column A               -> start of a part-number block, followed by
        "Planned release" row          -> written by the planner
        "Planed order" row             -> demand (input)
        "Cummulative" row              -> written by the planner (stock - backlog)
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

DAY_NAMES = ["Pon", "Uto", "Sre", "Cet", "Pet", "Sub", "Ned"]
FIRST_COL = 3  # column C


@dataclass
class PlanInput:
    path: Path
    pn_names: list[str]
    demand: np.ndarray          # (pn, days)
    capacity: np.ndarray        # (days,)
    day_names: list[str]
    week_of_day: np.ndarray     # 1-based week number per day
    release_rows: list[int]     # sheet rows to write results back into
    cumulative_rows: list[int]


def _num(v) -> float:
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _text(v) -> str:
    return v.strip() if isinstance(v, str) else ""


def read_input(path: str | Path) -> PlanInput:
    path = Path(path)
    ws = load_workbook(path, data_only=True).worksheets[0]
    max_row, max_col = ws.max_row, ws.max_column

    day_row = next(
        (r for r in range(1, min(30, max_row) + 1) if _text(ws.cell(r, FIRST_COL).value) in DAY_NAMES),
        None,
    )
    if day_row is None:
        raise ValueError(f"{path.name}: no row with day names (Pon, Uto, ...) in column C")
    day_names = []
    for c in range(FIRST_COL, max_col + 1):
        v = _text(ws.cell(day_row, c).value)
        if v not in DAY_NAMES:
            break
        day_names.append(v)
    days = len(day_names)

    cap_row = next(
        (r for r in range(1, day_row) if _text(ws.cell(r, FIRST_COL).value).lower() == "kapacitet"),
        None,
    )
    if cap_row is None:
        raise ValueError(f"{path.name}: no 'Kapacitet' row above the day names")
    capacity = np.array([_num(ws.cell(cap_row + 1, FIRST_COL + d).value) for d in range(days)])

    pn_names, demand, release_rows, cumulative_rows = [], [], [], []
    for r in range(day_row + 1, max_row + 1):
        name = _text(ws.cell(r, 1).value)
        if not name.startswith("PN-"):
            continue
        rows = {}
        for k in range(1, 4):
            label = _text(ws.cell(r + k, 1).value).lower()
            if "release" in label:
                rows["release"] = r + k
            elif "order" in label:
                rows["order"] = r + k
            elif "cum" in label:
                rows["cum"] = r + k
        if "order" not in rows:
            raise ValueError(f"{path.name}: {name} (row {r}) has no 'Planed order' row")
        pn_names.append(name)
        demand.append([_num(ws.cell(rows["order"], FIRST_COL + d).value) for d in range(days)])
        release_rows.append(rows.get("release"))
        cumulative_rows.append(rows.get("cum"))
    if not pn_names:
        raise ValueError(f"{path.name}: no PN-... rows found")

    # Week number: a new week starts at every Monday after the first day.
    week, weeks = 1, []
    for d, name in enumerate(day_names):
        if d > 0 and name == "Pon":
            week += 1
        weeks.append(week)

    return PlanInput(
        path=path,
        pn_names=pn_names,
        demand=np.array(demand),
        capacity=capacity,
        day_names=day_names,
        week_of_day=np.array(weeks),
        release_rows=release_rows,
        cumulative_rows=cumulative_rows,
    )


def _write_table(ws, header, rows, title=None):
    if title:
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([])
    ws.append(header)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = ws.cell(ws.max_row - len(rows) + 1, 2)


def write_report(inp: PlanInput, result, buckets: np.ndarray, out_path: str | Path) -> Path:
    """Copy the input workbook, fill in the plan, and add analysis sheets."""
    out_path = Path(out_path)
    wb = load_workbook(inp.path)
    ws = wb.worksheets[0]
    ws.title = "Plan"
    net = result.inventory - result.backlog
    for p in range(len(inp.pn_names)):
        for d in range(len(inp.day_names)):
            if inp.release_rows[p]:
                ws.cell(inp.release_rows[p], FIRST_COL + d, int(round(result.release[p, d])))
            if inp.cumulative_rows[p]:
                ws.cell(inp.cumulative_rows[p], FIRST_COL + d, int(round(net[p, d])))

    produced = result.release.sum(axis=0)
    dem = inp.demand.sum(axis=0)
    inv = result.inventory.sum(axis=0)
    bl = result.backlog.sum(axis=0)

    daily = []
    for d in range(len(inp.day_names)):
        util = produced[d] / inp.capacity[d] * 100 if inp.capacity[d] > 0 else None
        status = "BACKLOG" if bl[d] > 0 else ("OK" if util is None or util <= 100 else "PREKORAČENJE")
        daily.append([d + 1, f"CW{inp.week_of_day[d]:02d}", inp.day_names[d], inp.capacity[d],
                      dem[d], produced[d], None if util is None else round(util, 1), inv[d], bl[d], status])
    _write_table(wb.create_sheet("Dnevni pregled"),
                 ["Dan", "Nedelja", "Dan u nedelji", "Kapacitet", "Tražnja", "Proizvodnja",
                  "Iskorišćenost (%)", "Zalihe", "Backlog", "Status"], daily)

    weekly = []
    for w in np.unique(inp.week_of_day):
        m = inp.week_of_day == w
        cap = inp.capacity[m].sum()
        weekly.append([f"CW{w:02d}", cap, dem[m].sum(), produced[m].sum(),
                       round(produced[m].sum() / cap * 100, 1) if cap else None,
                       inv[m][-1], bl[m][-1]])
    _write_table(wb.create_sheet("Nedeljni pregled"),
                 ["Nedelja", "Kapacitet", "Tražnja", "Proizvodnja", "Iskorišćenost (%)",
                  "Zalihe na kraju", "Backlog na kraju"], weekly)

    bl_rows = [[name, *result.backlog[p].round().astype(int).tolist()] for p, name in enumerate(inp.pn_names)]
    _write_table(wb.create_sheet("Backlog"),
                 ["PN", *[f"D{d + 1:02d}" for d in range(len(inp.day_names))]], bl_rows)

    early_rows = []
    for p, name in enumerate(inp.pn_names):
        b = buckets[p]
        total = b.sum()
        early_rows.append([name, inp.demand[p].sum(), result.release[p].sum(), *b.round().astype(int).tolist(),
                           round(b[1:4].sum() / total * 100, 1) if total else 0.0])
    _write_table(wb.create_sheet("Proizvodnja unapred"),
                 ["PN", "Tražnja", "Proizvodnja", "Isti dan", "1 dan ranije", "2 dana ranije",
                  "3+ dana ranije", "Kasni (backlog)", "% unapred"], early_rows)

    for sheet in wb.worksheets[1:]:
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 14
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
